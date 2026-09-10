"""Weekly ticket-report export independent from the CLI adapter."""

from __future__ import annotations

import json
import os
import re
import time
from dataclasses import dataclass
from datetime import datetime
from importlib.resources import files
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .intraservice.client import IntraserviceClient
from .intraservice.parsing import TicketCard, parse_ticket_card
from .settings import PROJECT_ROOT, IntraserviceSettings

REPORT_HEADERS = (
    "Заявка",
    "Статус Б24",
    "Статус SD",
    "Тип ТП",
    "Описание",
    "Партнер",
    "Заказчик",
    "Текущий статус/решение",
    "Последнее обновление",
    "Исполнитель",
)

ProgressReporter = Callable[[str], None]
TICKET_TITLE_PREFIX_RE = re.compile(r"^\s*\d+\.\s*")
BRACKET_VALUE_RE = re.compile(r"\[([^\]]+)\]")
UPDATE_REQUEST_RE = re.compile(
    r"^Запрос обновления\s+\d+(?:[.-]\d+)*$", re.IGNORECASE
)
LATIN_LOOKALIKE_TO_CYRILLIC = str.maketrans(
    {
        "a": "а",
        "b": "в",
        "c": "с",
        "e": "е",
        "k": "к",
        "m": "м",
        "o": "о",
        "p": "р",
        "t": "т",
        "x": "х",
        "y": "у",
    }
)


def _normalize_organization(value: str) -> str:
    return " ".join(value.replace("«", '"').replace("»", '"').split()).casefold()


def _normalize_support_type(value: str) -> str:
    return " ".join(value.split()).casefold()


def _normalize_status(value: str) -> str:
    return " ".join(value.split()).casefold()


def _normalize_customer(value: str) -> str:
    return " ".join(value.split()).casefold().translate(LATIN_LOOKALIKE_TO_CYRILLIC)


def _private_config_path(env_name: str, filename: str) -> Path:
    raw_path = os.getenv(env_name)
    if raw_path and raw_path.strip():
        path = Path(raw_path.strip()).expanduser()
        return path if path.is_absolute() else PROJECT_ROOT / path
    return PROJECT_ROOT / ".local" / filename


def _load_private_json(env_name: str, filename: str, missing_value: object) -> object:
    source = _private_config_path(env_name, filename)
    try:
        return json.loads(source.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return missing_value


def load_partner_aliases() -> dict[str, str]:
    """Load the local creator-organization to partner mapping."""
    payload = _load_private_json("PARTNER_ALIASES_PATH", "partner_aliases.json", {})
    if not isinstance(payload, dict) or not all(
        isinstance(key, str) and isinstance(value, str) for key, value in payload.items()
    ):
        raise ValueError("partner_aliases.json должен содержать JSON-объект строковых алиасов.")
    return {_normalize_organization(key): value.strip() for key, value in payload.items()}


def load_support_type_aliases() -> dict[str, str]:
    """Load raw support-service names and their report-friendly aliases."""
    source = files("ispano").joinpath("support_type_aliases.json")
    payload = json.loads(source.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not all(
        isinstance(key, str) and isinstance(value, str) for key, value in payload.items()
    ):
        raise ValueError(
            "support_type_aliases.json должен содержать JSON-объект строковых алиасов."
        )
    return {_normalize_support_type(key): value.strip() for key, value in payload.items()}


def load_status_aliases() -> dict[str, str]:
    """Load raw SD status names and their report-friendly aliases."""
    source = files("ispano").joinpath("status_aliases.json")
    payload = json.loads(source.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not all(
        isinstance(key, str) and isinstance(value, str) for key, value in payload.items()
    ):
        raise ValueError("status_aliases.json должен содержать JSON-объект строковых алиасов.")
    return {_normalize_status(key): value.strip() for key, value in payload.items()}


def load_customer_names() -> dict[str, str]:
    """Load local allowed customer names keyed by their case-insensitive form."""
    payload = _load_private_json("CUSTOMER_NAMES_PATH", "customer_names.json", [])
    if not isinstance(payload, list) or not all(isinstance(value, str) for value in payload):
        raise ValueError("customer_names.json должен содержать JSON-массив строк.")
    return {_normalize_customer(value): value.strip() for value in payload if value.strip()}


def _parse_title_fields(
    title: str | None, customer_names: dict[str, str]
) -> tuple[str | None, str | None]:
    if not title:
        return None, None

    subject = TICKET_TITLE_PREFIX_RE.sub("", title, count=1).strip()
    customer = None
    for match in BRACKET_VALUE_RE.finditer(subject):
        customer = _resolve_customer(match.group(1), customer_names)
        if customer:
            break

    description_subject = BRACKET_VALUE_RE.sub("", subject).strip()
    description = (
        description_subject if UPDATE_REQUEST_RE.fullmatch(description_subject) else None
    )
    return description, customer


def _resolve_customer(value: str, customer_names: dict[str, str]) -> str | None:
    normalized = _normalize_customer(value)
    if not normalized:
        return None

    exact = customer_names.get(normalized)
    if exact:
        return exact

    prefix_matches = {
        canonical
        for alias, canonical in customer_names.items()
        if alias.startswith(normalized)
    }
    return next(iter(prefix_matches)) if len(prefix_matches) == 1 else None


@dataclass(frozen=True, slots=True)
class TicketReportRow:
    ticket_id: int
    status: str | None
    support_type: str | None
    partner: str | None
    last_updated_at: datetime | None
    description: str | None = None
    customer: str | None = None

    @classmethod
    def from_card(
        cls,
        ticket_id: int,
        card: TicketCard,
        aliases: dict[str, str],
        support_type_aliases: dict[str, str] | None = None,
        customer_names: dict[str, str] | None = None,
        status_aliases: dict[str, str] | None = None,
    ) -> "TicketReportRow":
        partner = None
        if card.creator_organization:
            partner = aliases.get(
                _normalize_organization(card.creator_organization),
                card.creator_organization,
            )
        support_type = card.support_type
        if support_type and support_type_aliases:
            support_type = support_type_aliases.get(
                _normalize_support_type(support_type), support_type
            )
        status = card.status
        if status and status_aliases:
            status = status_aliases.get(_normalize_status(status), status)
        description, customer = _parse_title_fields(card.title, customer_names or {})
        return cls(
            ticket_id,
            status,
            support_type,
            partner,
            card.last_updated_at,
            description,
            customer,
        )

    def values(self) -> tuple[object, ...]:
        return (
            self.ticket_id,
            "",
            self.status or "",
            self.support_type or "",
            self.description or "",
            self.partner or "",
            self.customer or "",
            "",
            self.last_updated_at,
            "",
        )


class TicketReportExporter:
    """Fetch report rows from IntraService ticket cards."""

    def __init__(self, settings: IntraserviceSettings, login: str, password: str) -> None:
        self.settings = settings
        self.login = login
        self.password = password

    def export(self, until_id: int, report: ProgressReporter = print) -> tuple[list[TicketReportRow], set[str]]:
        aliases = load_partner_aliases()
        support_type_aliases = load_support_type_aliases()
        status_aliases = load_status_aliases()
        customer_names = load_customer_names()
        unknown_organizations: set[str] = set()
        with IntraserviceClient(self.settings, self.login, self.password) as client:
            ticket_ids = client.list_ticket_ids_descending(until_id)
            report(f"Найдено тикетов для отчёта: {len(ticket_ids)}")
            rows: list[TicketReportRow] = []
            for index, ticket_id in enumerate(ticket_ids, 1):
                report(f"[{index}/{len(ticket_ids)}] Тикет {ticket_id}...")
                card = parse_ticket_card(client.get_ticket_page(ticket_id))
                row = TicketReportRow.from_card(
                    ticket_id,
                    card,
                    aliases,
                    support_type_aliases,
                    customer_names,
                    status_aliases,
                )
                rows.append(row)
                if (
                    card.creator_organization
                    and _normalize_organization(card.creator_organization) not in aliases
                ):
                    unknown_organizations.add(card.creator_organization)
                time.sleep(self.settings.request_delay)
        return rows, unknown_organizations


def write_ticket_report(rows: list[TicketReportRow], output_dir: Path) -> Path:
    """Write a copy-ready single-sheet XLSX report."""
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"tickets_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тикеты"
    sheet.append(REPORT_HEADERS)
    for row in rows:
        sheet.append(row.values())

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{max(len(rows) + 1, 1)}"
    for row in sheet.iter_rows(min_row=2):
        row[8].number_format = "dd.mm.yyyy"
        row[7].alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in zip("ABCDEFGHIJ", (12, 18, 22, 30, 35, 20, 22, 45, 20, 20), strict=True):
        sheet.column_dimensions[column].width = width
    workbook.save(path)
    return path
